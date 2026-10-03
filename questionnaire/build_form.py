"""
build_form.py
=============
Generates household_survey.xlsx, an XLSForm for KoboToolbox (or ODK).
Every question name here is read by survey_to_schema.py, so change the
two files together. Run:  python build_form.py
"""
import sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font

HERE = Path(__file__).parent

# ---------------------------------------------------------------------------
# Appliance catalogue shown in the form. Names match schema.py exactly.
# ---------------------------------------------------------------------------
ALWAYS_ON = ["refrigerator", "chest_freezer", "wifi_router",
             "electric_fence_energiser", "cctv_system"]
APPLIANCES = [
    ("refrigerator", "Fridge"), ("chest_freezer", "Chest freezer"),
    ("wifi_router", "Wi-Fi router"), ("electric_fence_energiser", "Electric fence energiser"),
    ("cctv_system", "CCTV system"),
    ("electric_kettle", "Electric kettle"), ("iron_box", "Iron box"),
    ("water_pump", "Water pump (tank / pressure)"),
    ("immersion_water_heater", "Instant shower or water heater"),
    ("solar_water_heater_pump", "Solar water heater pump"),
    ("electric_hotplate", "Electric hotplate / coil cooker"), ("induction_cooker", "Induction cooker"),
    ("electric_pressure_cooker", "Electric pressure cooker"), ("rice_cooker", "Rice cooker"),
    ("microwave", "Microwave"), ("blender", "Blender"), ("toaster", "Toaster"),
    ("electric_oven", "Electric oven"),
    ("television", "Television"), ("dstv_decoder", "Decoder (DStv, GOtv, StarTimes)"),
    ("laptop", "Laptop"), ("desktop_computer", "Desktop computer"),
    ("gaming_console", "Gaming console"), ("bluetooth_speaker", "Speaker / home theatre"),
    ("smartphone_charger", "Phone charging"), ("tablet_charger", "Tablet charging"),
    ("power_bank_charging", "Power bank charging"),
    ("washing_machine", "Washing machine"), ("vacuum_cleaner", "Vacuum cleaner"),
    ("ceiling_fan", "Ceiling fan"), ("standing_fan", "Standing / table fan"),
    ("air_conditioner", "Air conditioner"),
    ("gate_motor", "Gate motor"), ("borehole_pump", "Borehole pump"),
    ("hair_dryer", "Hair dryer"), ("electric_shaver", "Electric shaver / clipper"),
    ("sewing_machine", "Sewing machine"), ("printer", "Printer"),
    ("other", "Something else (type the name)"),
]

# Fallback values for downstream use when the form's own data is missing or
# a named appliance's label couldn't be read (label_watts/label_volts/
# label_amps all left blank). survey_to_schema.py should use these, in that
# situation only -- schema.py's own SURVEY NOTE for rated_power_w is
# explicit that a literature default must never override a readable label.
# Both tables are transcribed from REFERENCE_HOUSEHOLD's own values in
# schema.py, which are themselves Nairobi-specific literature/typical
# defaults for exactly this purpose. Not used for "other" appliances --
# those have no fixed identity to look up, so their power and occupancy
# dependence are asked directly (label_watts/volts/amps same as any
# appliance; needs_occupancy via appliance_other_home).
#
# The appliance repeat has no repeat_count limit -- an interviewer can add
# as many entries of the same appliance_type as the household actually has
# (e.g. three televisions with three different usage patterns), and
# survey_to_schema.py must give each one its own unique schema.py name
# (television, television_2, television_3, ...). Only television,
# electric_kettle and laptop have an explicit "_2" entry below, matching
# REFERENCE_HOUSEHOLD's illustrative example of one slightly-different
# second unit (e.g. a smaller bedroom TV); there is no "_3" entry for any
# appliance, and there never needs to be one. A device's typical wattage
# and whether it needs occupancy don't actually depend on whether it's the
# household's 1st, 2nd or 5th one of that type, so survey_to_schema.py's
# lookup should be: try the exact name first (catches the modelled "_2"
# case), and if that misses, strip any trailing "_N" and look up the base
# name instead. That handles any number of duplicates correctly without
# the table ever needing a "_3", "_4", ... entry.
#
# needs_occupancy is never asked for named appliances at all (same reasoning
# as calibrate_tou.APPLIANCE_RESTART_DELAY_MIN: it is a fixed property of
# what the appliance IS, not a household behaviour) -- this table is its
# only source for named appliances.
APPLIANCE_POWER_FALLBACK_W = {
    "refrigerator": 150, "chest_freezer": 120, "wifi_router": 12,
    "electric_fence_energiser": 25, "cctv_system": 30,
    "electric_kettle": 2000, "electric_kettle_2": 2000, "iron_box": 1200,
    "water_pump": 750, "immersion_water_heater": 3000, "solar_water_heater_pump": 50,
    "electric_hotplate": 1500, "induction_cooker": 2000, "electric_pressure_cooker": 800,
    "rice_cooker": 500, "microwave": 900, "blender": 350, "toaster": 800,
    "electric_oven": 2000,
    "television": 80, "television_2": 60, "dstv_decoder": 18,
    "laptop": 45, "laptop_2": 45, "desktop_computer": 150, "gaming_console": 150,
    "bluetooth_speaker": 10,
    "smartphone_charger": 10, "tablet_charger": 18, "power_bank_charging": 10,
    "washing_machine": 500, "vacuum_cleaner": 1000,
    "ceiling_fan": 60, "standing_fan": 50, "air_conditioner": 1500,
    "gate_motor": 200, "borehole_pump": 1500,
    "hair_dryer": 1500, "electric_shaver": 15, "sewing_machine": 100, "printer": 15,
}

APPLIANCE_NEEDS_OCCUPANCY = {
    "refrigerator": False, "chest_freezer": False, "wifi_router": False,
    "electric_fence_energiser": False, "cctv_system": False,
    "electric_kettle": True, "electric_kettle_2": True, "iron_box": True,
    "water_pump": False, "immersion_water_heater": True, "solar_water_heater_pump": False,
    "electric_hotplate": True, "induction_cooker": True, "electric_pressure_cooker": True,
    "rice_cooker": True, "microwave": True, "blender": True, "toaster": True,
    "electric_oven": True,
    "television": True, "television_2": True, "dstv_decoder": True,
    "laptop": True, "laptop_2": True, "desktop_computer": True, "gaming_console": True,
    "bluetooth_speaker": True,
    "smartphone_charger": False, "tablet_charger": False, "power_bank_charging": False,
    "washing_machine": True, "vacuum_cleaner": True,
    "ceiling_fan": False, "standing_fan": True, "air_conditioner": True,
    "gate_motor": False, "borehole_pump": False,
    "hair_dryer": True, "electric_shaver": True, "sewing_machine": True, "printer": True,
}

ROOMS = [
    ("living_room", "Living room / sitting room"), ("dining_area", "Dining area"),
    ("master_bedroom", "Main bedroom"), ("children_bedroom", "Children's bedroom"),
    ("bedroom", "Other bedroom"), ("kitchen", "Kitchen"), ("bathroom", "Bathroom / toilet"),
    ("staircase_corridor", "Corridor / stairs"), ("store_room", "Store room"),
    ("outside_security", "Outside / security lights"), ("other", "Other room"),
]

def hours(end=False):
    rng = range(1, 25) if end else range(0, 24)
    return [(str(h), f"{h:02d}:00" + (" (midnight)" if h in (0, 24) else "")) for h in rng]

CHOICES = {
    "yes_no": [("yes", "Yes"), ("no", "No")],
    "hour": hours(),
    "hour_end": hours(end=True),
    "hour_or_home": [("home_all_day", "Home all day, does not go out")] + hours(),
    "peak": [("spread", "No particular hour, spread evenly")] + hours(),
    "often": [("always", "Always"), ("usually", "Usually"),
              ("sometimes", "Sometimes"), ("rarely", "Rarely")],
    # keys match calibrate_tou.WEEKDAY_FREQUENCY / WEEKEND_FREQUENCY
    "weekday_frequency": [("every_weekday", "Every weekday (all 5)"), ("most_weekdays", "Most weekdays (about 4)"),
                ("few_times_a_week", "2 or 3 weekdays"), ("once_a_week", "About once a week"),
                ("rarely", "Less than once a week"), ("never", "Never on weekdays")],
    "weekend_frequency": [("both_days", "Both Saturday and Sunday"), ("usually_one_day", "Usually one of the two days"),
                ("occasionally", "Some weekends, not most"), ("never", "Never on weekends")],
    "appliance": APPLIANCES,
    "room": ROOMS,
    "bulb_type": [("LED", "LED"), ("CFL", "Energy saver (CFL, spiral)"),
                  ("incandescent", "Old-style filament bulb"), ("fluorescent", "Tube light")],
    "security_mode": [("dusk_to_dawn", "On all night, every night"), ("timer", "On a timer"),
                      ("manual", "Switched on and off by hand"), ("motion", "Motion sensor only")],
    "dwelling": [("flat", "Flat / apartment"), ("bungalow", "Bungalow"),
                 ("maisonette", "Maisonette / townhouse"), ("other", "Other")],
    "phase": [("single", "Single-phase"), ("three", "Three-phase"), ("dont_know", "Don't know")],
    "metering": [("prepaid", "Prepaid (tokens)"), ("postpaid", "Postpaid (monthly bill)")],
    "tariff": [("domestic", "Domestic"), ("lifeline", "Lifeline"), ("dont_know", "Don't know")],
    "reliability": [("good", "Good: rare outages"), ("fair", "Fair: a few outages a month"),
                    ("poor", "Poor: outages most weeks")],
    "backup": [("none", "None"), ("generator", "Generator"),
               ("inverter_battery", "Inverter and battery, no panels"),
               ("solar_only", "Solar panels, no battery"), ("solar_battery", "Solar panels and battery")],
    "orientation": [("north_facing", "North"), ("south_facing", "South"), ("east_facing", "East"),
                    ("west_facing", "West"), ("flat", "Flat roof")],
    "shading": [("none", "None"), ("minimal", "Minimal"), ("moderate", "Moderate"), ("severe", "Severe")],
    "roof_type": [("iron_sheet", "Iron sheet"), ("concrete", "Concrete slab"),
                  ("clay_tile", "Clay / concrete tile"), ("other", "Other")],
    "mounting": [("flush", "Flush on roof"), ("elevated", "Raised frame"), ("ground", "Ground")],
    "willing": [("yes", "Yes"), ("maybe", "Maybe"), ("no", "No")],
}

S = []  # survey rows
def q(type_, name="", label="", hint="", required="", constraint="", cmsg="",
      relevant="", appearance="", calculation="", repeat_count="", default="", parameters=""):
    S.append(dict(type=type_, name=name, label=label, hint=hint, required=required,
                  constraint=constraint, constraint_message=cmsg, relevant=relevant,
                  appearance=appearance, calculation=calculation,
                  repeat_count=repeat_count, default=default,
                  parameters=parameters or ("max-pixels=1024" if type_ == "image" else "")))

def window_repeat(name, label, relevant, day):
    """One repeat of usage windows; fields feed calibrate_tou.build_raw_shape().
    Field names are prefixed with the repeat name because XLSForm names must be unique."""
    p = name
    q("begin_repeat", name, label, relevant=relevant, appearance="field-list")
    q("select_one hour", f"{p}_start", f"{day}: from what time is it used?",
      hint="If it is used at more than one time of day, add an entry for each one.",
      required="yes", appearance="minimal")
    q("select_one hour_end", f"{p}_end", "Until what time?",
      hint="The hour it stops. 06:00 to 09:00 means up to 08:59. Times past midnight are allowed.",
      required="yes", constraint=f". != ${{{p}_start}}", cmsg="Start and end cannot be the same hour.",
      appearance="minimal")
    q("select_one peak", f"{p}_peak", "Within that time, which hour is it most likely to be on?",
      required="yes", appearance="minimal")
    q("select_one often", f"{p}_often",
      "How likely is it to actually be used at this particular time of day: always, usually, "
      "sometimes or rarely?",
      hint="If this is the only time of day you added, choose 'Always'. If you added more than one "
           "time of day, compare them: for example, evenings 'Always' and mornings 'Sometimes', if "
           "evenings are more reliable.",
      required="yes", appearance="horizontal")
    q("end_repeat")

SESSIONS_HINT = ("For example: a kettle boiled once in the morning and once in the evening, every day "
                 "it's used, is 2. A kettle boiled either in the morning or the evening, but not both, "
                 "is 1.")

def sessions_q(name, label, relevant):
    """Separate switch-ons per day of use (calibrate_tou interview note i/j)."""
    q("integer", name, label, hint=SESSIONS_HINT, relevant=relevant, required="yes",
      default="1", constraint=". >= 1 and . <= 20", cmsg="Between 1 and 20.")

def people_repeat(name, label, relevant, day, who):
    p = name
    q("begin_repeat", name, label, relevant=relevant, appearance="field-list")
    q("select_one hour", f"{p}_start", f"{day}: from what time?", required="yes", appearance="minimal")
    q("select_one hour_end", f"{p}_end", "Until what time?", required="yes",
      constraint=f". != ${{{p}_start}}", cmsg="Start and end cannot be the same hour.",
      appearance="minimal")
    q("integer", f"{p}_people", who, required="yes",
      constraint=". >= 0 and . <= ${number_of_residents}", cmsg="Cannot be more than the number of residents.")
    q("end_repeat")

# Whole-household occupancy (Section C) is asked per resident category, not
# as pre-aggregated hourly headcounts -- nobody naturally thinks of their
# household as "3 people home between 07:00 and 09:00". Every category is
# gated the same way, for both weekday and weekend: ask whether they go out
# at all on a normal day of that type before asking for a window. Even
# categories Section B defines by usually going out (working adults, school
# children) only describe a general pattern, not a guarantee for every
# single day -- remote work, a day off, an irregular schedule can all mean
# "home the whole day" on what's otherwise a normal weekday. Gating
# universally means that answer is always available, for every category,
# on both day types.
# survey_to_schema.py sums count x presence across all six categories per
# hour to build occupancy_weekday/occupancy_weekend -- the interviewer
# never aggregates headcounts by hand.
RESIDENT_CATEGORIES = [
    ("adults_working",     "the adults who go out to work or college"),
    ("adults_non_working", "the adults who are usually at home during the day"),
    ("household_helpers",  "the live-in house helps"),
    ("school_children",    "the children who go to school or day care"),
    ("young_children",     "the young children who stay at home during the day"),
    ("elderly",            "the older residents who are mostly at home"),
]

def category_presence(cat, who, day, relevant):
    """One category's typical weekday/weekend presence pattern. day must be
    the full word ("Weekday"/"Weekend") -- it is used both in the visible
    question text and, lowercased, as the field name prefix. Whether they
    go out at all and when they leave are a single question -- away_start's
    choice list includes "Home all day" directly, so there is no separate
    yes/no gate before it. away_end only appears once a real leave time is
    chosen."""
    p = f"{cat}_{day.lower()}"
    q("select_one hour_or_home", f"{p}_away_start",
      f"{day}: what time do {who} usually go out? Choose 'Home all day' if they do not go out.",
      relevant=relevant, required="yes", appearance="minimal")
    out = f"{relevant} and ${{{p}_away_start}} != 'home_all_day'"
    q("select_one hour_end", f"{p}_away_end", "Until what time?",
      relevant=out, required="yes",
      constraint=f". != ${{{p}_away_start}}", cmsg="Start and end cannot be the same hour.",
      appearance="minimal")

# ---------------------------------------------------------------- A. Start
q("start", "start"); q("end", "end")
q("begin_group", "group_start", "A. Consent and household details", appearance="field-list")
q("acknowledge", "consent",
  "Has the respondent read (or been read) the information sheet, had their questions answered, and signed the consent form?",
  required="yes")
q("text", "household_id", "Household ID", hint="Assigned by the researcher, for example H003.",
  required="yes", constraint="regex(., '^H[0-9]{3}$')", cmsg="Use H followed by three digits, e.g. H003.")
q("date", "survey_date", "Interview date", required="yes", default="today()")
q("text", "sub_county", "Sub-county", required="yes")
q("text", "county", "County", required="yes")
q("geopoint", "gps", "Record the GPS location standing at the house.",
  hint="Stored for the solar data only. Reports give the sub-county, never the exact point.")
q("select_one dwelling", "dwelling", "Type of house", required="yes")
q("end_group")

# ---------------------------------------------------------------- B. Composition
q("begin_group", "group_people", "B. Who lives here", appearance="field-list")
q("integer", "number_of_residents", "Including yourself, how many people normally sleep in this house on most nights?",
  hint="Include live-in house helps. Leave out visitors staying less than a month.",
  required="yes", constraint=". >= 1 and . <= 15", cmsg="Between 1 and 15.")
q("note", "residents_note", "Of these ${number_of_residents} people, how many are... (count each person once)")
for name, label in [
    ("adults_working", "adults who go out to work or college on most weekdays?"),
    ("adults_non_working", "adults who are usually at home during the day?"),
    ("household_helpers", "live-in house helps?"),
    ("school_children", "children who go to school or day care?"),
    ("young_children", "young children who stay at home during the day?"),
    ("elderly", "older people (about 65 and over) who are mostly at home?"),
]:
    q("integer", name, label, required="yes", default="0", constraint=". >= 0")
q("calculate", "resident_count_sum", calculation="${adults_working} + ${adults_non_working} + ${household_helpers} + ${school_children} + ${young_children} + ${elderly}")
q("note", "resident_count_check", "These add up to ${resident_count_sum} but there are ${number_of_residents} residents. Go back and correct.",
  relevant="${resident_count_sum} != ${number_of_residents}")
q("end_group")

# ---------------------------------------------------------------- C. Occupancy
q("begin_group", "group_occupancy", "C. When people are usually home", appearance="field-list")
q("note", "occupancy_note",
  "A few questions for each group of residents from Section B (only shown if the "
  "household has any). We'll ask about weekdays first, then whether weekends differ.")
for cat, who in RESIDENT_CATEGORIES:
    cond = f"${{{cat}}} > 0"
    category_presence(cat, who, "Weekday", cond)
    q("select_one yes_no", f"{cat}_weekend_same",
      f"Is this the same on a normal Saturday and Sunday for {who}?",
      relevant=cond, required="yes")
    category_presence(cat, who, "Weekend", f"{cond} and ${{{cat}_weekend_same}} = 'no'")
q("end_group")

# ---------------------------------------------------------------- D. Rooms and lights
q("begin_group", "group_rooms", "D. Rooms and lights (walk through the house room by room)")
q("begin_repeat", "room", "Room", appearance="field-list")
q("select_one room", "room_type", "Which room is this?", required="yes")
q("text", "room_other", "What is this room called?", relevant="${room_type} = 'other'", required="yes")
q("integer", "bulbs", "How many light bulbs or tubes are in this room?", required="yes", constraint=". >= 0")
q("decimal", "bulb_watts", "Wattage printed on the bulb (W)",
  hint="Read it off the bulb. If the bulbs differ, record the most common and note the others below.",
  relevant="${bulbs} > 0", constraint=". > 0 and . <= 200", cmsg="Between 1 and 200 W. Check units.")
q("select_one bulb_type", "bulb_type", "Bulb type", relevant="${bulbs} > 0", required="yes")
q("image", "bulb_photo", "Photo of a bulb label (optional)", relevant="${bulbs} > 0")
# outside lights
q("select_one security_mode", "security_mode", "How are the outside lights switched?",
  relevant="${room_type} = 'outside_security' and ${bulbs} > 0", required="yes")
q("select_one hour", "security_light_on", "What time do they usually go on?",
  relevant="${security_mode} = 'timer' or ${security_mode} = 'manual'", required="yes",
  appearance="minimal")
q("select_one hour_end", "security_light_off", "What time do they usually go off?",
  relevant="${security_mode} = 'timer' or ${security_mode} = 'manual'", required="yes",
  constraint=". != ${security_light_on}", cmsg="Start and end cannot be the same hour.",
  appearance="minimal")
# indoor rooms
# room_weekday/room_weekend feed each bulb's mean_duration_min AND
# std_duration_min -- no separate duration/variability question is asked.
# Each repeat entry's own (end - start) is that visit's mean;
# std_duration_min is derived downstream as ~25% of that mean, matching
# calibrate_tou.py's own documented fallback for when std can't be reliably
# asked directly. Not computed across repeat entries -- a room's morning
# vs. evening blocks are structurally different sessions, not repeated
# samples of the same event, and there are usually too few of them for a
# meaningful standard deviation.
inside = "${room_type} != 'outside_security'"
q("note", "room_presence_note",
  "When is someone in this room on a normal weekday, and how many people? "
  "(Awake time only, not sleeping.) "
  "This also tells us when the light gets used -- we work out from location and date "
  "when it's dark enough to need it, so there's no need to ask about the light separately.",
  relevant=inside)
people_repeat("room_weekday", "Weekday: time in this room", inside, "Weekday",
              "How many people are usually in this room then?")
q("select_one yes_no", "room_weekend_same", "Is the room used at the same times at the weekend?", relevant=inside, required="yes")
people_repeat("room_weekend", "Weekend: time in this room", f"{inside} and ${{room_weekend_same}} = 'no'", "Weekend",
              "How many people are usually in this room then?")
q("text", "room_notes", "Notes on this room (optional)")
q("end_repeat")
q("end_group")

# ---------------------------------------------------------------- E. Appliances
q("begin_group", "group_appliances", "E. Appliances")
q("note", "appliance_note",
  "Add one entry for each appliance, reading the label as you go. If two of the same appliance are used "
  "at different times (for example a sitting-room TV and a bedroom TV), add them as two entries. "
  "Before finishing, ask: anything that charges, heats water, pumps water, cooks, or runs all night?")
q("begin_repeat", "appliance", "Appliance", appearance="field-list")
q("select_one appliance", "appliance_type", "Which appliance?", required="yes", appearance="minimal autocomplete")
q("text", "appliance_other", "What is it?", relevant="${appliance_type} = 'other'", required="yes")
q("select_one yes_no", "appliance_other_always_on",
  "Does it run continuously all the time, like a fridge or router, rather than being switched on and off?",
  relevant="${appliance_type} = 'other'", required="yes")
q("select_one yes_no", "appliance_other_home", "Does it only run when someone is at home?",
  relevant="${appliance_type} = 'other' and ${appliance_other_always_on} = 'no'", required="yes")
q("integer", "appliance_count", "How many of these are used together in the same way?", default="1",
  required="yes", constraint=". >= 1 and . <= 20")
q("decimal", "label_watts", "Rated power on the label (W)",
  hint="Look for W or kW on the rating plate (1.5 kW = 1500 W). If it shows only volts and amps, enter them below. "
       "If there is no label at all, leave this, volts and amps all blank -- a typical value for this "
       "appliance will be used instead.",
  constraint=". > 0 and . <= 10000", cmsg="Between 1 and 10000 W.")
q("decimal", "label_volts", "Volts on the label", relevant="not(${label_watts} > 0)")
q("decimal", "label_amps", "Amps on the label", relevant="not(${label_watts} > 0) and ${label_volts} > 0")
q("decimal", "standby_watts", "Standby power on the label, if printed (W)", constraint=". >= 0")
q("image", "label_photo", "Photo of the rating label")
# label_watts/volts/amps may all be left blank (no readable label).
# survey_to_schema.py must then fall back to
# APPLIANCE_POWER_FALLBACK_W[appliance_type] for named appliances --
# schema.py's validator requires rated_power_w > 0 whenever count > 0, so a
# blank label can never be allowed to reach schema.py as a missing value.
# For "other" appliances with no label, there is no literature default to
# fall back to; survey_to_schema.py should flag that case for the
# researcher rather than guessing.
q("calculate", "always_on",
  calculation="if(" + " or ".join(f"${{appliance_type}} = '{a}'" for a in ALWAYS_ON) +
  " or ${appliance_other_always_on} = 'yes', 'yes', 'no')")
use = "${always_on} = 'no'"
q("integer", "duration_typical",
  "The last time it was used, about how many minutes was it on?",
  hint="For charging: time plugged in. For a washing machine: one full wash. For a fan: until it was switched off.",
  relevant=use, required="yes", constraint=". >= 1 and . <= 1440")
q("integer", "duration_shortest", "What is the shortest time it is usually on (minutes)?", relevant=use,
  constraint=". >= 1 and . <= ${duration_typical}", cmsg="Should not be longer than the typical time.")
q("integer", "duration_longest", "And the longest (minutes)?", relevant=use,
  constraint=". >= ${duration_typical} and . <= 1440", cmsg="Should not be shorter than the typical time.")
q("select_one weekday_frequency", "weekday_frequency",
  "On how many weekdays (Monday to Friday) in a normal week is it used, even just once?",
  hint="This is about which days, not how many times in a day -- that's the next question.",
  relevant=use, required="yes")
window_repeat("appliance_weekday", "Weekday time of use", f"{use} and ${{weekday_frequency}} != 'never'", "Weekday")
sessions_q("appliance_weekday_sessions",
           "On a day it's used, how many separate times does it get switched on?",
           f"{use} and ${{weekday_frequency}} != 'never'")
q("select_one weekend_frequency", "weekend_frequency",
  "And in a normal weekend, on how many of the two days is it used, even just once?",
  relevant=use, required="yes")
q("select_one yes_no", "weekend_same_as_weekday", "On weekend days it is used, is it used at the same times as on weekdays?",
  relevant=f"{use} and ${{weekend_frequency}} != 'never' and ${{weekday_frequency}} != 'never'", required="yes")
window_repeat("appliance_weekend", "Weekend time of use",
              f"{use} and ${{weekend_frequency}} != 'never' and (${{weekday_frequency}} = 'never' or ${{weekend_same_as_weekday}} = 'no')", "Weekend")
sessions_q("appliance_weekend_sessions",
           "On a weekend day it's used, how many separate times does it get switched on?",
           f"{use} and ${{weekend_frequency}} != 'never'")
q("text", "appliance_notes", "Notes (optional)")
q("end_repeat")
q("end_group")

# ---------------------------------------------------------------- F. Grid
# schema.py's grid block (Block 6) also has import_tariff_kes_per_kwh and
# monthly_fixed_charge_kes -- neither is asked here. Both are Objective 3
# (MILP sizing) inputs, not currently in scope; like export_tariff_kes_per_kwh
# (schema.py: fixed at 0, no Kenyan feed-in tariff exists) they are standard
# KPLC rates that barely vary household to household, not something an
# interview needs to elicit per household. tariff (below) gives the rate
# category (domestic/lifeline); the actual KES/kWh figure should come from
# a small, separately maintained rate table (current KPLC rates), the same
# "physical/structural fact, not interview data" reasoning as
# calibrate_tou.APPLIANCE_RESTART_DELAY_MIN and this file's
# APPLIANCE_POWER_FALLBACK_W/APPLIANCE_NEEDS_OCCUPANCY. Revisit when
# Objective 3 work starts.
q("begin_group", "group_grid", "F. Electricity supply and backup", appearance="field-list")
q("select_one yes_no", "grid_connected", "Is the house connected to Kenya Power?", required="yes")
g = "${grid_connected} = 'yes'"
q("select_one phase", "phase", "Is the supply single-phase or three-phase?",
  hint="Three-phase meters have three incoming live wires or a 3-phase label.", relevant=g, required="yes")
q("select_one metering", "metering", "Prepaid tokens or a monthly bill?", relevant=g, required="yes")
q("select_one tariff", "tariff", "Tariff shown on the bill or token SMS", relevant=g, required="yes")
q("integer", "monthly_spend", "In a normal month, about how much do you spend on electricity, bills or tokens (KES)?",
  relevant=g, required="yes", constraint=". >= 0")
q("decimal", "token_kwh", "Last token: how many kWh did it give? (from the SMS)", relevant="${metering} = 'prepaid'")
q("integer", "token_kes", "And for how many shillings?", relevant="${token_kwh} > 0")
q("select_one reliability", "reliability", "How reliable is the power supply here?", relevant=g, required="yes")
q("decimal", "blackout_hours", "In a normal week, about how many hours in total is the power off?",
  relevant=g, required="yes", constraint=". >= 0 and . <= 168")
q("select_one backup", "backup", "Do you have any backup power?", required="yes")
q("integer", "solar_panel_count", "How many solar panels?", relevant="${backup} = 'solar_only' or ${backup} = 'solar_battery'", required="yes")
q("decimal", "solar_panel_watts", "Watts per panel (from the panel label)", relevant="${solar_panel_count} > 0", required="yes")
bat = "${backup} = 'inverter_battery' or ${backup} = 'solar_battery'"
q("decimal", "battery_volts", "Battery voltage (V, from label)", relevant=bat, required="yes")
q("decimal", "battery_amp_hours", "Battery capacity (Ah, from label)", relevant=bat, required="yes")
q("integer", "battery_units", "Number of batteries", relevant=bat, required="yes", default="1")
q("decimal", "backup_kilowatts", "Inverter or generator rated output (kW, from label)",
  relevant="${backup} != 'none'", required="yes")
q("text", "backup_loads", "What does the backup power when the grid is off?", relevant="${backup} != 'none'")
q("end_group")

# ---------------------------------------------------------------- G. Site
q("begin_group", "group_site", "G. Roof and site (interviewer, from the walk-around)", appearance="field-list")
q("decimal", "roof_length", "Usable roof length (m)", required="yes", constraint=". > 0")
q("decimal", "roof_width", "Usable roof width (m)", required="yes", constraint=". > 0")
q("select_one orientation", "orientation", "Which way does the best roof face?", required="yes")
q("integer", "tilt", "Roof tilt (degrees, estimate)", required="yes", constraint=". >= 0 and . <= 60")
q("select_one shading", "shading", "Shading from trees or buildings", required="yes")
q("select_one roof_type", "roof_type", "Roof material", required="yes")
q("select_one mounting", "mounting", "Likely panel mounting", required="yes")
q("decimal", "cable_length_meters", "Cable run from roof to a likely inverter spot (m)", required="yes")
q("end_group")

# ---------------------------------------------------------------- H. Close
q("begin_group", "group_close", "H. Follow-up", appearance="field-list")
q("select_one willing", "logger_willing",
  "Would you be willing to host a power logger on your main board for one week, including a weekend? "
  "A qualified installer would fit it; it only records how much power the house uses.", required="yes")
q("text", "routine_change", "Any weeks in the next two months when your routine will be unusual (travel, guests, holidays)?")
q("note", "readback",
  "Read back the busiest electricity times you recorded and ask if they sound right. Correct any entries before saving.")
q("text", "interviewer_notes", "Interviewer notes", appearance="multiline")
q("end_group")

# ---------------------------------------------------------------- write xlsx
def build(out=HERE / "household_survey.xlsx"):
    wb = Workbook()
    ws = wb.active
    assert ws is not None  # a freshly created Workbook always has one default sheet
    ws.title = "survey"
    cols = ["type", "name", "label", "hint", "required", "constraint", "constraint_message",
            "relevant", "appearance", "calculation", "repeat_count", "default", "parameters"]
    ws.append(cols)
    for r in S:
        ws.append([r[c] for c in cols])
    ch = wb.create_sheet("choices"); ch.append(["list_name", "name", "label"])
    for lst, items in CHOICES.items():
        for n, l in items:
            ch.append([lst, n, l])
    st = wb.create_sheet("settings")
    st.append(["form_title", "form_id", "version", "default_language", "style"])
    # style "pages": the Kobo web form shows each section (A to H) as its own page.
    st.append(["Household energy use interview", "household_energy_v1", "1", "English", "pages"])
    for sh in wb.worksheets:
        for c in sh[1]:
            c.font = Font(bold=True)
    wb.save(out)
    print(f"wrote {out}  ({len(S)} survey rows)")

if __name__ == "__main__":
    build()
